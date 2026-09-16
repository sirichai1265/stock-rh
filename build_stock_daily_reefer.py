# -*- coding: utf-8 -*-
"""
Stock Daily Reefer Report
=========================
Input : stock/staying .xls  +  booking .xls
Output: Stock_Daily_<M-D>.xlsx   (Control / BookingRaw / Summary, formula-driven)
        Stock_Daily_<M-D>.pptx   (8-slide dashboard)
        Stock_Daily_<M-D>.model.json  (computed model, for verification)

Run:  python build_stock_daily_reefer.py "9-10-STAYING - Copy.xls" "9-10-BKG+PD - Copy.xls"
(defaults to those two files in the script folder if no args given)

Recalc of the .xlsx and PNG export of the .pptx are done by companion
PowerShell scripts that drive Excel / PowerPoint COM.
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# config
# --------------------------------------------------------------------------- #
FOLDER = Path(__file__).resolve().parent

# raw code -> display type   (same mapping for stock Size/Type and booking columns)
BKG_TYPE_ORDER = ["GP22", "GP42", "GP45", "RE22", "RE45", "UT22", "UT42", "PC22", "PC42"]
CODE2DISP = {
    "GP22": "20'GP", "GP42": "40'GP", "GP45": "40'HC",
    "RE22": "20'RE", "RE45": "40'RH",
    "UT22": "20'OT", "UT42": "40'OT",
    "PC22": "20'FR", "PC42": "40'FR",
}
# stock Size/Type codes use the same 2-letter families in reversed order (22RE ...)
SIZE2DISP = {
    "22GP": "20'GP", "42GP": "40'GP", "45GP": "40'HC",
    "22RE": "20'RE", "45RE": "40'RH",
    "22UT": "20'OT", "42UT": "40'OT",
    "22PC": "20'FR", "42PC": "40'FR",
}
DISP_ORDER = ["20'GP", "40'GP", "40'HC", "20'RE", "40'RH", "20'OT", "40'OT", "20'FR", "40'FR"]
RE_DISP = {"20'RE", "40'RH"}
OTFR_DISP = {"20'OT", "40'OT", "20'FR", "40'FR"}

# location blocks --------------------------------------------------------------
# BKK side: row-major fill of lane1 (A-K) then lane2 (M-W)
#   (BKK25 , BKK01+BKK04)
#   (BKK27 , BKK02)
#   (empty , LCH55)
# LCH side: lane3 (Y-AI) stacked
BKK_SEQUENCE = [
    ("BKK25", ["BKK25"]),
    ("BKK01+BKK04", ["BKK01", "BKK04"]),
    ("BKK27", ["BKK27"]),
    ("BKK02", ["BKK02"]),
    (None, None),                       # empty slot (lane1, 3rd row-block)
    ("LCH55", ["LCH55"]),
]
LCH_SEQUENCE = [
    ("LCH27", ["LCH27"]),
    ("LCH28", ["LCH28"]),
    ("LCHY5", ["LCHY5"]),
]

# colours
NAVY = "0C2340"
SUBTITLE_BG = "E7ECF3"
STOCK_BG = "F3F5F8"
AV_BG = "FFF6C8"          # yellow
RE_FONT = "1F4E9C"        # blue
OTFR_FONT = "7A4B12"      # brown
NEG_FONT = "C22A2A"       # red

BLOCK_H = 13
LANE_STARTS = {1: 1, 2: 13, 3: 25}      # A , M , Y
GAP_COLS = [12, 24]                      # L , X
ROWBLOCK_TOPS = [4, 17, 30]

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# --------------------------------------------------------------------------- #
# step 1 - read
# --------------------------------------------------------------------------- #
def load_stock(path: Path) -> dict[tuple[str, str], int]:
    df = pd.read_excel(path, header=0)
    df = df[df["Location"].notna() & df["Size/Type"].notna()]
    df = df[df["Full/Empty"].astype(str).str.strip().str.upper() == "E"]
    df["loc"] = df["Location"].astype(str).str.strip().str.upper()
    df["disp"] = df["Size/Type"].astype(str).str.strip().str.upper().map(SIZE2DISP)
    df = df[df["disp"].notna()]
    agg: dict[tuple[str, str], int] = {}
    for (loc, disp), n in df.groupby(["loc", "disp"]).size().items():
        agg[(loc, disp)] = int(n)
    return agg


def load_booking(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path, header=0)
    if "Pickup" not in df.columns:
        df = pd.read_excel(path, header=1)
    if "Pickup" not in df.columns:
        raise SystemExit("booking file: no 'Pickup' column at header 0 or 1")
    df = df[df["Pickup"].notna() & df["TRAN DT"].notna()].copy()
    df["loc"] = df["Pickup"].astype(str).str.strip().str.upper()
    df["loc"] = df["loc"].replace({"BKK04": "BKK01"})          # merge BKK04 -> BKK01
    df["date"] = pd.to_datetime(df["TRAN DT"].astype(float).astype("int64").astype(str),
                                format="%Y%m%d")
    out = pd.DataFrame({"Location": df["loc"].values, "Date": df["date"].values})
    for code in BKG_TYPE_ORDER:
        col = code if code in df.columns else None
        vals = pd.to_numeric(df[col], errors="coerce").fillna(0).astype(int) if col else 0
        out[CODE2DISP[code]] = vals.values if col else 0
    return out.reset_index(drop=True)


# --------------------------------------------------------------------------- #
# step 2/4 - pure-python model (mirrors the Excel formulas, drives the pptx)
# --------------------------------------------------------------------------- #
def week_frame(today: dt.date):
    mon1 = today - dt.timedelta(days=today.weekday())
    weeks = [(mon1 + dt.timedelta(days=7 * i), mon1 + dt.timedelta(days=7 * i + 5)) for i in range(4)]
    tomorrow = today + dt.timedelta(days=1)
    iso = [d[0].isocalendar()[1] for d in weeks]
    return weeks, tomorrow, iso


def build_model(stock_agg, bkg: pd.DataFrame, today: dt.date):
    weeks, tomorrow, iso = week_frame(today)
    sat1 = weeks[0][1]
    bkg = bkg.copy()
    bkg["d"] = pd.to_datetime(bkg["Date"]).dt.date

    def block_for(codes):
        sub = bkg[bkg["Location"].isin([c.upper() for c in codes])]
        types = {}
        for disp in DISP_ORDER:
            stock = sum(stock_agg.get((c.upper(), disp), 0) for c in codes)
            col = sub[disp] if disp in sub.columns else pd.Series([], dtype=int)
            pend = int(sub.loc[sub["d"] < today, disp].sum()) if len(sub) else 0
            t0 = int(sub.loc[sub["d"] == today, disp].sum()) if len(sub) else 0
            rest = int(sub.loc[(sub["d"] > today) & (sub["d"] <= sat1), disp].sum()) if len(sub) else 0
            wk = []
            for i, (a, b) in enumerate(weeks):
                if i == 0:
                    wk.append(t0 + rest)
                else:
                    wk.append(int(sub.loc[(sub["d"] >= a) & (sub["d"] <= b), disp].sum()) if len(sub) else 0)
            av, run = [], stock - pend
            for i in range(4):
                run -= wk[i]
                av.append(run)
            types[disp] = dict(stock=stock, pending=pend, today=t0, wk1rest=rest,
                               wk=wk, av=av)
        return types

    locs = []
    for label, codes in BKK_SEQUENCE + LCH_SEQUENCE:
        if label is None:
            continue
        zone = "LCH" if (label.startswith("LCH") and label != "LCH55") else "BKK"
        locs.append(dict(code=label, zone=zone, types=block_for(codes)))

    return dict(
        date=today.isoformat(),
        weeks=[[a.isoformat(), b.isoformat()] for a, b in weeks],
        iso_weeks=iso,
        tomorrow=tomorrow.isoformat(),
        locations=locs,
    )


# --------------------------------------------------------------------------- #
# step 3 - Excel
# --------------------------------------------------------------------------- #
def mn(ref: str) -> str:
    return f'CHOOSE(MONTH({ref}),"' + '","'.join(MONTHS) + '")'


def range_label(s: str, e: str) -> str:
    return (f'IF(MONTH({s})=MONTH({e}),DAY({s})&"-"&DAY({e})&" "&{mn(e)},'
            f'DAY({s})&" "&{mn(s)}&"-"&DAY({e})&" "&{mn(e)})')


def build_excel(model, stock_agg, bkg: pd.DataFrame, out: Path):
    wb = Workbook()
    F = Font(name="Calibri", size=11)
    FB = Font(name="Calibri", size=11, bold=True)
    thin = Side(style="thin", color="C9D2DF")
    med = Side(style="medium", color="0C2340")

    # ---- Control ----------------------------------------------------------
    ws = wb.active
    ws.title = "Control"
    ws["A1"] = "Report date"
    ws["B1"] = "=TODAY()"
    ws["B1"].number_format = "yyyy-mm-dd"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "WK#", "Monday", "Saturday", "ISO week"
    ws["B4"] = "=$B$1-(WEEKDAY($B$1,2)-1)"
    ws["C4"] = "=$B$4+5"
    for r in (5, 6, 7):
        ws[f"B{r}"] = f"=$B${r-1}+7"
        ws[f"C{r}"] = f"=$B${r}+5"
    for r in (4, 5, 6, 7):
        ws[f"A{r}"] = f"=_xlfn.ISOWEEKNUM($B${r})"
        ws[f"D{r}"] = f"=_xlfn.ISOWEEKNUM($B${r})"
        ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = "yyyy-mm-dd"
    ws["A9"] = "Rest of WK1 start"
    ws["B9"] = "=$B$1+1"
    ws["B9"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 18
    for L in ("B", "C", "D"):
        ws.column_dimensions[L].width = 13
    for row in ws.iter_rows():
        for c in row:
            c.font = F

    # ---- BookingRaw -----------------------------------------------------
    wr = wb.create_sheet("BookingRaw")
    wr.append(["Location", "Date"] + DISP_ORDER)
    for _, rr in bkg.iterrows():
        wr.append([rr["Location"], pd.to_datetime(rr["Date"]).to_pydatetime()]
                  + [int(rr[d]) for d in DISP_ORDER])
    for row in wr.iter_rows():
        for c in row:
            c.font = F
        if row[0].row > 1:
            row[1].number_format = "yyyy-mm-dd"
    for i, w in enumerate([16, 12] + [7] * 9):
        wr.column_dimensions[get_column_letter(i + 1)].width = w

    # ---- Summary ------------------------------------------------------
    sm = wb.create_sheet("Summary")
    last_col = LANE_STARTS[3] + 10                      # AI
    for i in range(1, last_col + 1):
        L = get_column_letter(i)
        if i in GAP_COLS:
            sm.column_dimensions[L].width = 2.3
        elif (i - 1) % 12 == 0:
            sm.column_dimensions[L].width = 5.5         # WK col
        elif (i - 2) % 12 == 0:
            sm.column_dimensions[L].width = 20          # label col
        else:
            sm.column_dimensions[L].width = 6.2         # type cols

    title = ('="Stock Daily - Reefer  (Stock vs Booking)   as of "&DAY(Control!$B$1)&" "&'
             + mn("Control!$B$1") + '&" "&YEAR(Control!$B$1)')
    sm.cell(1, 1, title).font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    sm.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    sub = ('="Report date  "&TEXT(Control!$B$1,"yyyy-mm-dd")&'
           '"      Legend:  RE = reefer (blue)   OT/FR = special (brown)   AV Balance = stock - cumulative booking (yellow; red = short)"')
    sm.cell(2, 1, sub).font = Font(name="Calibri", size=10, italic=True, color="333333")
    sm.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    for cc in range(1, last_col + 1):
        sm.cell(1, cc).fill = PatternFill("solid", fgColor=NAVY)
        sm.cell(2, cc).fill = PatternFill("solid", fgColor=SUBTITLE_BG)
    for lane, txt in ((1, "BKK"), (2, "BKK"), (3, "LCH")):
        c = sm.cell(3, LANE_STARTS[lane], txt)
        c.font = Font(name="Calibri", size=11, bold=True, color=NAVY)

    layout = []          # (label, codes, lane, top_row)
    for idx, (label, codes) in enumerate(BKK_SEQUENCE):
        lane = 1 if idx % 2 == 0 else 2
        top = ROWBLOCK_TOPS[idx // 2]
        if label is not None:
            layout.append((label, codes, lane, top))
    for idx, (label, codes) in enumerate(LCH_SEQUENCE):
        layout.append((label, codes, 3, ROWBLOCK_TOPS[idx]))

    mt = {loc["code"]: loc for loc in model["locations"]}

    for label, codes, lane, top in layout:
        c0 = LANE_STARTS[lane]           # WK col (1-idx)
        cN = c0 + 1                       # label col
        cT0 = c0 + 2                      # first type col
        wkL, nmL = get_column_letter(c0), get_column_letter(cN)
        loc_key = codes[0].upper() if len(codes) == 1 else "BKK01"

        r_title, r_head, r_stock, r_pend = top, top + 1, top + 2, top + 3
        r_t0, r_rest, r_av1 = top + 4, top + 5, top + 6
        r_w2, r_av2, r_w3, r_av3, r_w4, r_av4 = [top + i for i in range(7, 13)]

        # title
        tc = sm.cell(r_title, c0, label)
        tc.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        tc.alignment = Alignment(horizontal="center")
        sm.merge_cells(start_row=r_title, start_column=c0, end_row=r_title, end_column=c0 + 10)
        for cc in range(c0, c0 + 11):
            sm.cell(r_title, cc).fill = PatternFill("solid", fgColor=NAVY)

        # header
        sm.cell(r_head, c0, "WK").font = FB
        for i, disp in enumerate(DISP_ORDER):
            hc = sm.cell(r_head, cT0 + i, disp)
            col = RE_FONT if disp in RE_DISP else OTFR_FONT if disp in OTFR_DISP else "0C2340"
            hc.font = Font(name="Calibri", size=11, bold=True, color=col)
            hc.alignment = Alignment(horizontal="center")

        # stock empty in yard  (hard snapshot, every cell filled)
        sm.cell(r_stock, cN, "Stock empty in yard").font = FB
        for i, disp in enumerate(DISP_ORDER):
            v = sum(stock_agg.get((c.upper(), disp), 0) for c in codes)
            sc = sm.cell(r_stock, cT0 + i, v)
            sc.font = F
            sc.alignment = Alignment(horizontal="center")
        for cc in range(c0, c0 + 11):
            sm.cell(r_stock, cc).fill = PatternFill("solid", fgColor=STOCK_BG)

        def sumifs(colL, extra):
            return (f'=SUMIFS(BookingRaw!${colL}$2:${colL}$600,'
                    f'BookingRaw!$A$2:$A$600,"{loc_key}",{extra})')

        # pending
        sm.cell(r_pend, c0, "=Control!$A$4").font = F
        sm.cell(r_pend, cN, "Booking Pending pick up").font = F
        # wk1 today
        sm.cell(r_t0, c0, "=Control!$A$4").font = F
        sm.cell(r_t0, cN, f'="Booking on "&DAY(Control!$B$1)&" "&{mn("Control!$B$1")}').font = F
        # wk1 rest
        sm.cell(r_rest, c0, "=Control!$A$4").font = F
        sm.cell(r_rest, cN,
                f'=IF(Control!$B$9=Control!$C$4,DAY(Control!$B$9)&" "&{mn("Control!$B$9")},'
                f'"Booking on "&{range_label("Control!$B$9","Control!$C$4")})').font = F
        # av1 label
        sm.cell(r_av1, cN,
                f'="AV Balance till "&DAY(Control!$C$4)&" "&{mn("Control!$C$4")}').font = FB
        # wk2..4
        for rr, wk_ctrl, monref, satref in (
            (r_w2, "$A$5", "Control!$B$5", "Control!$C$5"),
            (r_w3, "$A$6", "Control!$B$6", "Control!$C$6"),
            (r_w4, "$A$7", "Control!$B$7", "Control!$C$7"),
        ):
            sm.cell(rr, c0, f"=Control!{wk_ctrl}").font = F
            sm.cell(rr, cN, f'="Booking on "&{range_label(monref, satref)}').font = F
        for rr, cref in ((r_av2, "$C$5"), (r_av3, "$C$6"), (r_av4, "$C$7")):
            sm.cell(rr, cN, f'="AV Balance till "&DAY(Control!{cref})&" "&{mn("Control!"+cref)}').font = FB

        # per-type formulas
        tmodel = mt[label]["types"]
        for i, disp in enumerate(DISP_ORDER):
            cc = cT0 + i
            TL = get_column_letter(cc)
            BL = get_column_letter(3 + i)              # BookingRaw C..K
            for rr, extra in (
                (r_pend, '"<"&Control!$B$1'),
                (r_t0, '"="&Control!$B$1'),
                (r_rest, '">="&Control!$B$9,BookingRaw!$B$2:$B$600,"<="&Control!$C$4'),
                (r_w2, '">="&Control!$B$5,BookingRaw!$B$2:$B$600,"<="&Control!$C$5'),
                (r_w3, '">="&Control!$B$6,BookingRaw!$B$2:$B$600,"<="&Control!$C$6'),
                (r_w4, '">="&Control!$B$7,BookingRaw!$B$2:$B$600,"<="&Control!$C$7'),
            ):
                e = sm.cell(rr, cc, sumifs(BL, f'BookingRaw!$B$2:$B$600,{extra}'))
                e.font = Font(name="Calibri", size=11,
                              color=RE_FONT if disp in RE_DISP else OTFR_FONT if disp in OTFR_DISP else "000000")
                e.alignment = Alignment(horizontal="center")
            # av rows
            sm.cell(r_av1, cc, f"={TL}{r_stock}-{TL}{r_pend}-{TL}{r_t0}-{TL}{r_rest}")
            sm.cell(r_av2, cc, f"={TL}{r_av1}-{TL}{r_w2}")
            sm.cell(r_av3, cc, f"={TL}{r_av2}-{TL}{r_w3}")
            sm.cell(r_av4, cc, f"={TL}{r_av3}-{TL}{r_w4}")
            av_vals = tmodel[disp]["av"]
            for k, rr in enumerate((r_av1, r_av2, r_av3, r_av4)):
                cell = sm.cell(rr, cc)
                neg = av_vals[k] < 0
                cell.font = Font(name="Calibri", size=11, bold=neg,
                                 color=NEG_FONT if neg else "000000")
                cell.alignment = Alignment(horizontal="center")

        # yellow fill on the 4 AV rows (full block width)
        for rr in (r_av1, r_av2, r_av3, r_av4):
            for cc in range(c0, c0 + 11):
                sm.cell(rr, cc).fill = PatternFill("solid", fgColor=AV_BG)

        # medium border around whole block
        r_bot = r_av4
        for rr in range(r_title, r_bot + 1):
            for cc in range(c0, c0 + 11):
                cell = sm.cell(rr, cc)
                cell.border = Border(
                    left=med if cc == c0 else thin,
                    right=med if cc == c0 + 10 else thin,
                    top=med if rr == r_title else thin,
                    bottom=med if rr == r_bot else thin,
                )

    # row heights / base font
    for r in range(1, ROWBLOCK_TOPS[-1] + BLOCK_H + 2):
        sm.row_dimensions[r].height = 13
    sm.sheet_view.showGridLines = False
    sm.freeze_panes = "A4"

    wb.save(out)


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def main():
    args = sys.argv[1:]
    stock_path = Path(args[0]) if len(args) > 0 else FOLDER / "9-10-STAYING - Copy.xls"
    bkg_path = Path(args[1]) if len(args) > 1 else FOLDER / "9-10-BKG+PD - Copy.xls"
    if not stock_path.is_absolute():
        stock_path = FOLDER / stock_path
    if not bkg_path.is_absolute():
        bkg_path = FOLDER / bkg_path

    today = dt.date.today()
    tag = f"{today.month}-{today.day}"

    stock_agg = load_stock(stock_path)
    bkg = load_booking(bkg_path)
    model = build_model(stock_agg, bkg, today)

    (FOLDER / f"Stock_Daily_{tag}.model.json").write_text(
        json.dumps(model, indent=2, ensure_ascii=False), encoding="utf-8")

    build_excel(model, stock_agg, bkg, FOLDER / f"Stock_Daily_{tag}.xlsx")
    print("xlsx written:", f"Stock_Daily_{tag}.xlsx")

    import build_pptx_reefer
    build_pptx_reefer.build(model, FOLDER / f"Stock_Daily_{tag}.pptx")
    print("pptx written:", f"Stock_Daily_{tag}.pptx")


if __name__ == "__main__":
    main()
