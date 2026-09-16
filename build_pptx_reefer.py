# -*- coding: utf-8 -*-
"""Stock Daily Reefer - 8-slide dashboard (python-pptx).  build(model, out_path)."""
from __future__ import annotations

import datetime as dt

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION,
                             XL_TICK_LABEL_POSITION)
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

NAVY = RGBColor(0x0C, 0x23, 0x40)
BKK = RGBColor(0xC8, 0x7A, 0x17)
LCH = RGBColor(0x0E, 0x7C, 0x86)
RED = RGBColor(0xC2, 0x2A, 0x2A)
GREEN = RGBColor(0x1E, 0x8E, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x2A, 0x33)
MUTE = RGBColor(0x5B, 0x6B, 0x7B)
CARD = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xD5, 0xDD, 0xE6)

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
EMU_W, EMU_H = Emu(12192000), Emu(6858000)      # 16:9
RE_TYPES = ["20'RE", "40'RH"]


def _md(iso):
    d = dt.date.fromisoformat(iso)
    return f"{d.day} {MONTHS[d.month - 1]}"


def _box(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    sp.adjustments[0] = 0.06
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for txt, size, bold, color in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = "Calibri"
    return tb


def _bg(slide, color):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); sp.shadow.inherit = False
    slide.shapes._spTree.remove(sp._element)
    slide.shapes._spTree.insert(2, sp._element)
    return sp


def _header(slide, title, sub):
    _box(slide, 0, 0, EMU_W, Emu(760000), fill=NAVY)
    _text(slide, Emu(430000), Emu(90000), Emu(9000000), Emu(360000),
          [(title, 24, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, Emu(430000), Emu(470000), Emu(11200000), Emu(230000),
          [(sub, 11, False, RGBColor(0xC7, 0xD2, 0xE0))])


def _sum_types(loc):
    """reefer totals for one location"""
    t = loc["types"]
    stock = sum(t[k]["stock"] for k in RE_TYPES)
    pend = sum(t[k]["pending"] for k in RE_TYPES)
    av4 = sum(t[k]["av"][3] for k in RE_TYPES)
    worst = min((t[k]["av"][i] for k in RE_TYPES for i in range(4)), default=0)
    first_neg = None
    for i in range(4):
        if any(t[k]["av"][i] < 0 for k in RE_TYPES):
            first_neg = i
            break
    return stock, pend, av4, worst, first_neg


def build(model, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = EMU_W, EMU_H
    blank = prs.slide_layouts[6]
    locs = model["locations"]
    iso = model["iso_weeks"]
    wk_lab = [f"WK{iso[0]}", f"WK{iso[1]}", f"WK{iso[2]}", f"WK{iso[3]}"]
    wk_span = [f"{_md(model['weeks'][i][0])}-{_md(model['weeks'][i][1])}" for i in range(4)]
    rdate = _md(model["date"]) + " " + model["date"][:4]

    bkk = [l for l in locs if l["zone"] == "BKK"]
    lch = [l for l in locs if l["zone"] == "LCH"]

    def zstock(zl):
        return sum(sum(l["types"][k]["stock"] for k in RE_TYPES) for l in zl)

    def zpend(zl):
        return sum(sum(l["types"][k]["pending"] for k in RE_TYPES) for l in zl)

    total_stock = zstock(bkk) + zstock(lch)
    total_pend = zpend(bkk) + zpend(lch)
    watch = [l for l in locs if _sum_types(l)[4] is not None]
    earliest = None
    for l in watch:
        fn = _sum_types(l)[4]
        if earliest is None or fn < earliest:
            earliest = fn

    # ---- slide 1 : title ------------------------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, NAVY)
    _text(s, Emu(900000), Emu(2100000), Emu(10400000), Emu(900000),
          [("Stock Daily – Reefer Report", 40, True, WHITE)])
    _text(s, Emu(900000), Emu(3050000), Emu(10400000), Emu(500000),
          [("Empty reefer stock vs outstanding bookings  ·  20'RE / 40'RH", 16, False,
            RGBColor(0xC7, 0xD2, 0xE0))])
    _text(s, Emu(900000), Emu(3750000), Emu(10400000), Emu(400000),
          [("Report date  " + rdate, 14, False, RGBColor(0x9F, 0xB1, 0xC6))])
    _box(s, Emu(900000), Emu(4500000), Emu(220000), Emu(220000), fill=BKK)
    _text(s, Emu(1200000), Emu(4470000), Emu(3000000), Emu(300000),
          [("BKK depots", 12, False, WHITE)])
    _box(s, Emu(2500000), Emu(4500000), Emu(220000), Emu(220000), fill=LCH)
    _text(s, Emu(2800000), Emu(4470000), Emu(3000000), Emu(300000),
          [("LCH depots", 12, False, WHITE)])

    # ---- slide 2 : KPI overview --------------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, WHITE)
    _header(s, "KPI Overview", "Reefer empties on hand against bookings · " + rdate)
    cards = [
        ("Total empty stock", str(total_stock), "20'RE + 40'RH, both zones", NAVY),
        ("Total pending pickup", str(total_pend), "bookings overdue (before today)",
         RED if total_pend else GREEN),
        ("Locations at risk", str(len(watch)), "AV balance turns negative within 4 wks",
         RED if watch else GREEN),
        ("Earliest shortage", wk_lab[earliest] if earliest is not None else "–",
         wk_span[earliest] if earliest is not None else "none in 4-week window",
         RED if earliest is not None else GREEN),
    ]
    cw, gap, x0, y0 = 2760000, 200000, 430000, 1050000
    for i, (lab, val, note, col) in enumerate(cards):
        x = x0 + i * (cw + gap)
        _box(s, x, y0, cw, 1850000, fill=CARD, line=LINE)
        _text(s, x + 150000, y0 + 130000, cw - 300000, 300000, [(lab, 11, True, MUTE)])
        _text(s, x + 150000, y0 + 470000, cw - 300000, 800000, [(val, 40, True, col)])
        _text(s, x + 150000, y0 + 1330000, cw - 300000, 460000, [(note, 9.5, False, MUTE)])

    # zone split table
    ty = y0 + 2150000
    _text(s, x0, ty, 6000000, 300000, [("Zone split", 14, True, NAVY)])
    ty += 380000
    rows = [
        ("", "Empty stock", "Pending pickup", "AV bal. end of WK" + str(iso[3])),
        ("BKK", str(zstock(bkk)), str(zpend(bkk)),
         str(sum(_sum_types(l)[2] for l in bkk))),
        ("LCH", str(zstock(lch)), str(zpend(lch)),
         str(sum(_sum_types(l)[2] for l in lch))),
    ]
    colw = [1500000, 3000000, 3000000, 3200000]
    for r, row in enumerate(rows):
        cx = x0
        for cix, cell in enumerate(row):
            head = r == 0
            _box(s, cx, ty, colw[cix], 360000,
                 fill=NAVY if head else (CARD if r % 2 else WHITE),
                 line=LINE)
            neg = (not head and cix == 3 and cell.lstrip("-").isdigit() and int(cell) < 0)
            _text(s, cx + 90000, ty + 40000, colw[cix] - 150000, 300000,
                  [(cell, 10.5, head or cix == 0, WHITE if head else (RED if neg else INK))])
            cx += colw[cix]
        ty += 360000

    # ---- slide 3 : bar chart stock by type BKK vs LCH ----------------
    s = prs.slides.add_slide(blank)
    _bg(s, WHITE)
    _header(s, "Empty Stock by Type – BKK vs LCH", "Reefer empties on hand · " + rdate)
    cd = CategoryChartData()
    cd.categories = RE_TYPES
    cd.add_series("BKK", tuple(sum(l["types"][k]["stock"] for l in bkk) for k in RE_TYPES))
    cd.add_series("LCH", tuple(sum(l["types"][k]["stock"] for l in lch) for k in RE_TYPES))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(900000), Emu(1050000),
                            Emu(10400000), Emu(5100000), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.plots[0].has_data_labels = True
    ch.series[0].format.fill.solid(); ch.series[0].format.fill.fore_color.rgb = BKK
    ch.series[1].format.fill.solid(); ch.series[1].format.fill.fore_color.rgb = LCH

    # ---- slide 4 : line chart AV balance trend -----------------------
    s = prs.slides.add_slide(blank)
    _bg(s, WHITE)
    _header(s, "AV Balance Trend – Weekly", "Empty stock minus cumulative booking, by zone · " + rdate)
    cd = CategoryChartData()
    cd.categories = [f"{wk_lab[i]}\n{wk_span[i]}" for i in range(4)]
    cd.add_series("BKK", tuple(sum(_sum_types(l)[2] if False else
                                   sum(l["types"][k]["av"][i] for k in RE_TYPES) for l in bkk)
                              for i in range(4)))
    cd.add_series("LCH", tuple(sum(sum(l["types"][k]["av"][i] for k in RE_TYPES) for l in lch)
                              for i in range(4)))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Emu(900000), Emu(1050000),
                            Emu(10400000), Emu(4700000), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    for ser, col in ((ch.series[0], BKK), (ch.series[1], LCH)):
        ser.format.line.color.rgb = col
        ser.format.line.width = Pt(2.25)
        ser.smooth = False
        m = ser.marker
        m.style = 8  # circle
        m.format.fill.solid(); m.format.fill.fore_color.rgb = col
        m.format.line.color.rgb = col
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(11)
    plot.data_labels.number_format = "0;-0"
    plot.data_labels.number_format_is_linked = False
    try:
        plot.data_labels.position = XL_LABEL_POSITION.ABOVE
    except Exception:
        pass
    ch.category_axis.tick_labels.font.size = Pt(10)
    ch.value_axis.tick_labels.font.size = Pt(10)
    try:
        ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    except Exception:
        pass
    _text(s, Emu(900000), Emu(5950000), Emu(10400000), Emu(500000),
          [("Values below zero = not enough empties to cover that week's bookings.", 10.5, False, MUTE)])

    # ---- slide 5 : locations to watch -------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, WHITE)
    _header(s, "Locations to Watch", "Where reefer AV balance turns negative in the 4-week window")
    y = 1100000
    if not watch:
        _box(s, 430000, y, 11330000, 950000, fill=CARD, line=GREEN, line_w=1.5)
        _text(s, 640000, y + 300000, 11000000, 500000,
              [("No depot goes short on 20'RE or 40'RH within the next 4 weeks.", 15, True, GREEN)])
    for l in watch:
        t = l["types"]
        stock, pend, av4, worst, fn = _sum_types(l)
        ch = 2500000
        _box(s, 430000, y, 11330000, ch, fill=CARD, line=RED, line_w=1.75)
        _box(s, 430000, y, 90000, ch, fill=RED)
        _text(s, 640000, y + 150000, 4000000, 420000, [(l["code"], 20, True, NAVY)])
        _text(s, 640000, y + 560000, 6000000, 320000,
              [(f"empty stock {stock}   ·   pending pickup {pend}", 11, False, MUTE)])
        _box(s, 9250000, y + 170000, 2300000, 380000, fill=RED)
        _text(s, 9250000, y + 195000, 2300000, 340000,
              [(f"first short  {wk_lab[fn]}", 11, True, WHITE)], align=PP_ALIGN.CENTER)
        ry = y + 980000
        for k in RE_TYPES:
            avs = t[k]["av"]
            nw = next((i for i in range(4) if avs[i] < 0), None)
            _text(s, 640000, ry, 1300000, 300000, [(k, 12, True, INK)])
            if nw is None:
                _text(s, 2000000, ry, 8000000, 300000,
                      [(f"covered all 4 weeks (ends {avs[3]:+d})", 11, False, GREEN)])
            else:
                runs = []
                for i in range(4):
                    runs.append((f"{wk_lab[i]} {avs[i]:+d}    ", 10.5, i == nw,
                                 RED if avs[i] < 0 else GREEN))
                _text(s, 2000000, ry, 9200000, 300000, [runs])
                short_need = -avs[3]
                _text(s, 640000, ry + 300000, 10500000, 300000,
                      [(f"     shortfall builds to {avs[3]:+d} by WK{iso[3]} "
                        f"— needs ~{short_need} more {k} at {l['code']} before {_md(model['weeks'][nw][0])}",
                        10, False, MUTE)])
            ry += 640000
        y += ch + 250000
    if watch:
        ok_codes = [l["code"] for l in locs if _sum_types(l)[4] is None
                    and (sum(l["types"][k]["stock"] for k in RE_TYPES) > 0
                         or sum(l["types"][k]["wk"][j] for k in RE_TYPES for j in range(4)) > 0)]
        extra = f"  ({', '.join(ok_codes)})" if ok_codes else ""
        _text(s, 640000, y + 120000, 11000000, 400000,
              [(f"All other depots stay covered on 20'RE and 40'RH through WK{iso[3]}.{extra}",
                11.5, False, GREEN)])

    # ---- slides 6 & 7 : location detail cards -----------------------
    for zone, zl, zcol in (("BKK", bkk, BKK), ("LCH", lch, LCH)):
        s = prs.slides.add_slide(blank)
        _bg(s, WHITE)
        _header(s, f"Location Detail – {zone}", "Stock / Pending / AV balance per depot · " + rdate)
        n = len(zl)
        cols = min(3, max(1, n))
        rows_n = (n + cols - 1) // cols
        cw = int((11330000 - (cols - 1) * 200000) / cols)
        chh = 1750000 if rows_n > 2 else 2300000
        for i, l in enumerate(zl):
            r, c = divmod(i, cols)
            x = 430000 + c * (cw + 200000)
            yy = 1050000 + r * (chh + 180000)
            stock, pend, av4, worst, fn = _sum_types(l)
            ok = fn is None
            _box(s, x, yy, cw, chh, fill=CARD, line=LINE)
            _box(s, x, yy, cw, 70000, fill=zcol)
            _text(s, x + 150000, yy + 150000, cw - 900000, 360000, [(l["code"], 15, True, NAVY)])
            _box(s, x + cw - 950000, yy + 150000, 800000, 300000, fill=GREEN if ok else RED)
            _text(s, x + cw - 950000, yy + 155000, 800000, 290000,
                  [("OK" if ok else "WATCH", 10, True, WHITE)], align=PP_ALIGN.CENTER)
            kv = [
                ("Empty stock", str(stock)),
                ("Pending pickup", str(pend)),
                (f"AV bal. end WK{iso[3]}", f"{av4:+d}"),
            ]
            ly = yy + 600000
            for lab, val in kv:
                _text(s, x + 170000, ly, cw - 1400000, 280000, [(lab, 10.5, False, MUTE)])
                neg = val.startswith("-")
                _text(s, x + cw - 1250000, ly, 1100000, 280000,
                      [(val, 12, True, RED if neg else INK)], align=PP_ALIGN.RIGHT)
                ly += 330000
            per = []
            for k in RE_TYPES:
                a = l["types"][k]["av"][3]
                per.append((f"{k} {a:+d}   ", 10, False, RED if a < 0 else GREEN))
            _text(s, x + 170000, ly + 40000, cw - 300000, 260000, [per])

    # ---- slide 8 : key takeaways -----------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, WHITE)
    _header(s, "Key Takeaways", "Auto-generated from today's stock and booking data")
    pts = []
    if watch:
        for l in watch:
            t = l["types"]
            for k in RE_TYPES:
                avs = t[k]["av"]
                nw = next((i for i in range(4) if avs[i] < 0), None)
                if nw is None:
                    continue
                need = -avs[3]
                pts.append((f"{l['code']} – {k}: goes short from {wk_lab[nw]} "
                            f"({wk_span[nw]}), reaching {avs[3]:+d} by WK{iso[3]}. "
                            f"Reposition ≥ {need} × {k} into {l['code']} before {_md(model['weeks'][nw][0])}.",
                            RED))
    else:
        pts.append(("All depots hold enough reefer empties to cover bookings through WK"
                    + str(iso[3]) + ".", GREEN))
    healthy = [l for l in locs if _sum_types(l)[4] is None
               and sum(l["types"][k]["stock"] for k in RE_TYPES) > 0]
    if healthy:
        pts.append(("Healthy: " + ", ".join(l["code"] for l in healthy)
                    + " – positive AV balance every week.", GREEN))
    zero_stock = [l["code"] for l in locs
                  if sum(l["types"][k]["stock"] for k in RE_TYPES) == 0
                  and sum(l["types"][k]["wk"][j] for k in RE_TYPES for j in range(4)) > 0]
    if zero_stock:
        pts.append(("Booking demand but zero empty stock on hand: " + ", ".join(zero_stock)
                    + " – every pickup there needs repositioned units.", NAVY))
    tot_book = sum(sum(l["types"][k]["wk"][j] for k in RE_TYPES for j in range(4)) for l in locs)
    pts.append((f"Reefer pipeline: {tot_book} units of 20'RE / 40'RH booked over the next 4 weeks, "
                f"plus {total_pend} already overdue — against {total_stock} empties on hand today.", NAVY))

    y = 1250000
    for txt, col in pts:
        _box(s, 470000, y + 55000, 200000, 200000, fill=col)
        _text(s, 880000, y - 30000, 10600000, 1100000, [(txt, 14, False, INK)], spacing=1.12)
        y += 620000 + 300000 * (len(txt) // 88)

    prs.save(str(out_path))


if __name__ == "__main__":
    import json, sys
    m = json.load(open(sys.argv[1], encoding="utf-8"))
    build(m, sys.argv[2])
