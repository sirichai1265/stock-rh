# -*- coding: utf-8 -*-
import re, sys, warnings
import openpyxl
from openpyxl.styles import Font
import formulas

F = sys.argv[1]
warnings.filterwarnings("ignore")
xl = formulas.ExcelModel().loads(F).finish()
sol = xl.calculate()

vals = {}
for k, v in sol.items():
    m = re.search(r"\]([^'!\]]+)'?!([A-Z]+\d+)$", k)
    if not m:
        continue
    sheet, coord = m.group(1).strip("'").upper(), m.group(2)
    try:
        val = v.value[0, 0]
    except Exception:
        val = getattr(v, "value", None)
    vals.setdefault(sheet, {})[coord] = val

summ = vals.get("TOTAL RH", {})
wb = openpyxl.load_workbook(F)
sm = wb["TOTAL RH"]
RED = Font(color="C62828", bold=True)
NORMAL = Font(bold=True)

# Find AV-Balance rows generically: any row whose label formula (col B or H,
# i.e. one column right of a "WK" header col) contains "AV Balance"
label_cols = []
for c in range(1, sm.max_column + 1):
    header = sm.cell(3, c).value
    if header == "WK":
        label_cols.append(c + 1)  # label sits immediately right of WK col

report = {}
neg_cells = []
for lbl_col in label_cols:
    base = lbl_col  # label col == base, values at base+1, base+2 (20'RE, 40'RH)
    for r in range(4, sm.max_row + 1):
        v = sm.cell(r, lbl_col).value
        if v and "AV Balance" in str(v):
            for off in (1, 2):
                col = openpyxl.utils.get_column_letter(base + off)
                cell = "%s%d" % (col, r)
                val = summ.get(cell)
                neg = isinstance(val, (int, float)) and val < 0
                sm[cell].font = RED if neg else NORMAL
                report[cell] = val
                if neg:
                    neg_cells.append(cell)

wb.calculation.fullCalcOnLoad = True
wb.save(F)
print("AV balances:", report)
print("negative:", neg_cells)
